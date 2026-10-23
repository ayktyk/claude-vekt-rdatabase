#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""isbu_excel.py — avukatın iş takip Excel'i (VEGA HUKUK .xlsx) → isbu-ofis SADE öneriler.

Avukat kararı (20.10.2026): isbu-ofis'e giden öneriler Excel satırı kadar sade olacak. Bu betik Excel'i tek
ölçü alır: DOSYALAR sayfasındaki her açık satır için isbu davasını bulur, isbu'da BOŞ olan alanları (esas no,
mahkeme, özet) Excel'den doldurmayı ve "Sıradaki İş" hücresini dosyanın "Beklenen" alanına yazmayı önerir; İŞLER
sayfasındaki açık işleri görev olarak önerir. Kuyrukta bekleyen eski önerileri de Excel'e göre sadeleştirir:
Excel'de olanlar Excel bilgisiyle yeniden yazılır; olmayanlar "Yeni dava (UYAP)" kalır (Excel yalnız Mücahit'le
ortak dosyaları tutar, isbu-ofis hepsini: Excel'de olmamak ret sebebi değildir — avukat 22.10.2026).

Tahmin yok: isbu davası yalnız (1) UYAP bağı, (2) mahkeme + esas no, (3) müvekkil adının BİREBİR eşleşmesi ile
bulunur; adla bulunan dava yalnız müvekkilin Excel'de tek açık dosyası varsa kullanılır. Birden çok aday ya da aynı
isbu davasına ikinci satır → belirsiz, öneri yok. Excel'de olup isbu'da bulunmayan satır için yeni dava önerisi
yalnız isbu'da ve kuyrukta benzer adlı hiçbir kayıt yoksa üretilir.

Anahtarlar satır numarasına değil içeriğe bağlıdır (müvekkil no + dosya no/kurum + konu; işlerde müvekkil + iş
metni özeti), satır eklenince/silinince başka kayda yazılmaz. Önceki çalıştırmadan kalan eski anahtarlı bekleyen
öneriler "Eski anahtar · reddedin" diye işaretlenir.

Kullanım:
  python scripts/isbu_excel.py plan [--gonder] [--cikti <klasör>] [--excel <xlsx>]
Çıktı: tmp/isbu-ofis/<tarih>-excel/ (gitignore'da; müvekkil verisi içerir). Gönderim yalnız --gonder ile.
"""
from __future__ import annotations

import argparse
import hashlib
import re
import sys
from datetime import date, datetime
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))
import isbu_esitle as ie  # noqa: E402
import isbu_sade as sade  # noqa: E402
import uyap_ortak as uo  # noqa: E402

REPO_ROOT = Path(__file__).resolve().parent.parent
TUZ = "sade-2026-10-20"
ETIKET = {"case_number": "esas no", "court_name": "mahkeme", "description": "özet", "buro_kodu": "büro kodu",
          "is_cmk_assignment": "CMK işareti", "title": "dava adı", "client_name": "müvekkil", "case_type": "dava türü",
          "custom_case_type": "özel tür", "start_date": "başlangıç", "contracted_fee": "ücret", "status": "durum"}


# --- Excel -------------------------------------------------------------------------------------------------------

def _sayfa_oku(ws, baslik_ipucu: str) -> list[dict]:
    """Başlık satırını ('Sıra' + ipucu sütunu) bulup altındaki dolu satırları sözlük yapar; tarihleri ISO güne çevirir."""
    basliklar, out = None, []
    for i, satir in enumerate(ws.iter_rows(values_only=True), start=1):
        hucreler = [str(c).strip() if c is not None else "" for c in satir]
        if basliklar is None:
            if hucreler and hucreler[0] == "Sıra" and baslik_ipucu in hucreler:
                basliklar = hucreler
            continue
        if not any(c is not None for c in satir):
            continue
        d = {}
        for b, c in zip(basliklar, satir):
            if not b:
                continue
            if isinstance(c, datetime):
                c = c.date().isoformat()
            elif isinstance(c, date):
                c = c.isoformat()
            d[b] = c
        d["_satir"] = i
        out.append(d)
    return out


def excel_oku(yol: Path) -> dict:
    import openpyxl
    wb = openpyxl.load_workbook(yol, read_only=True, data_only=True)
    try:
        dosyalar = _sayfa_oku(wb["DOSYALAR"], "Müvekkil") if "DOSYALAR" in wb.sheetnames else []
        isler = _sayfa_oku(wb["İŞLER"], "Yapılacak İş") if "İŞLER" in wb.sheetnames else []
    finally:
        wb.close()
    return {"dosyalar": dosyalar, "isler": isler}


def _s(d: dict, ad: str) -> str:
    v = d.get(ad)
    return "" if v is None else str(v).strip()


# --- eşleştirme --------------------------------------------------------------------------------------------------

def _kelimeler(ad) -> list[str]:
    return [w for w in re.split(r"[^a-z0-9]+", uo.tr(ad)) if w]


def _slug(metin, sinir: int = 30) -> str:
    return re.sub(r"[^a-z0-9]+", "-", uo.tr(metin)).strip("-")[:sinir]


def _kisa_hash(*parcalar) -> str:
    return hashlib.sha1("|".join(uo.tr(p) for p in parcalar).encode("utf-8")).hexdigest()[:16]


def ad_eslesir(excel_ad, isbu_ad, isbu_baslik="") -> bool:
    """Birebir ad eşleşmesi: Excel müvekkil adının kelime dizisi isbu müvekkil adınınkiyle aynı (büyük/küçük harf
    ve Türkçe karakter farkı yok sayılır). isbu müvekkil adı boşsa başlıkla karşılaştırılır."""
    kel = _kelimeler(excel_ad)
    return bool(kel) and kel == _kelimeler(isbu_ad or isbu_baslik)


def ad_benzer(excel_ad, isbu_ad, isbu_baslik="") -> bool:
    """Gevşek benzerlik (yalnız 'yeni dava önerme' frenidir): ortak bir anlamlı kelime (3+ harf) yeter."""
    kel = {w for w in _kelimeler(excel_ad) if len(w) >= 3}
    return bool(kel & {w for w in _kelimeler(isbu_ad) + _kelimeler(isbu_baslik) if len(w) >= 3})


def dava_bul(satir: dict, davalar: list[dict]) -> tuple[list[dict], str]:
    """(adaylar, yol). yol: uyap | esas | ad."""
    mahkeme, no = _s(satir, "Mahkeme / Kurum"), _s(satir, "Dosya No")
    if mahkeme and no:
        ref = ie.uyap_ref(mahkeme, no)
        aday = [d for d in davalar if d.get("externalRef") == ref]
        if aday:
            return aday, "uyap"
        aday = [d for d in davalar if d.get("caseNumber") and d.get("courtName")
                and uo.ayni_dosya(d["courtName"], d["caseNumber"], mahkeme, no)]
        if aday:
            return aday, "esas"
    ad = _s(satir, "Müvekkil")
    return [d for d in davalar if ad_eslesir(ad, d.get("clientName"), d.get("title"))], "ad"


def bekleyen_bul(p: dict, dosyalar: list[dict], dava: dict | None = None) -> list[dict]:
    """Kuyruktaki önerinin Excel satırı: mahkeme + esas no varsa yalnız onunla; yoksa müvekkil adıyla ve ancak
    müvekkilin Excel'de tek açık dosyası varsa (aksi halde boş: tahmin yok)."""
    pl = p.get("payload") or {}
    mahkeme, no = pl.get("court_name") or "", pl.get("case_number") or ""
    if mahkeme and no:
        return [s for s in dosyalar if _s(s, "Mahkeme / Kurum") and _s(s, "Dosya No")
                and uo.ayni_dosya(mahkeme, no, _s(s, "Mahkeme / Kurum"), _s(s, "Dosya No"))]
    ad = (dava or {}).get("clientName") or pl.get("client_name") or ""
    aday = [s for s in dosyalar if ad_eslesir(_s(s, "Müvekkil"), ad, (dava or {}).get("title") or pl.get("title") or "")]
    if len(aday) == 1 and sum(1 for s in dosyalar if _kelimeler(_s(s, "Müvekkil")) == _kelimeler(_s(aday[0], "Müvekkil"))) == 1:
        return aday
    return []


# --- öneriler ----------------------------------------------------------------------------------------------------

def _baslik(satir: dict, dosyalar: list[dict]) -> str:
    """Dava adı: 'Müvekkil – Konu'; aynı müvekkil + konu birden çok satırsa mahkeme eklenir."""
    ad, konu = _s(satir, "Müvekkil"), _s(satir, "Konu")
    ayni = [s for s in dosyalar if _kelimeler(_s(s, "Müvekkil")) == _kelimeler(ad) and _s(s, "Konu") == konu]
    baslik = f"{ad} – {konu}" if konu else ad
    if len(ayni) > 1 and _s(satir, "Mahkeme / Kurum"):
        baslik += f" ({_s(satir, 'Mahkeme / Kurum')} {_s(satir, 'Dosya No')})".replace(" )", ")")
    return baslik[:500]


def _ozet(satir: dict) -> str:
    return sade.ozet_satiri(_s(satir, "Müvekkil"), _s(satir, "Karşı Taraf"), _s(satir, "Mahkeme / Kurum"),
                            _s(satir, "Dosya No"), _s(satir, "Konu"), _s(satir, "Aşama"))


def _gerekce(satir: dict) -> str:
    return sade.gerekce(_s(satir, "Sıradaki İş"), f"ortak Excel satır {satir['_satir']}")


def _satir_kimligi(satir: dict) -> str:
    """Satırın kalıcı anahtarı (satır numarasına bağlı DEĞİL): müvekkil no (yoksa müvekkil adı özeti) + dosya no
    (yoksa kurum adı) + konu. Aynı müvekkilin numarasız iki dosyası konuyla ayrışır."""
    mno = _s(satir, "Müvekkil No") or _kisa_hash(_s(satir, "Müvekkil"))
    ek = uo.dosya_no(_s(satir, "Dosya No")) if _s(satir, "Dosya No") else _slug(_s(satir, "Mahkeme / Kurum"))
    return ":".join(x for x in (mno, ek, _slug(_s(satir, "Konu"), 20)) if x)


def _dis_kimlik(satir: dict) -> str:
    mahkeme, no = _s(satir, "Mahkeme / Kurum"), _s(satir, "Dosya No")
    if mahkeme and no:
        return ie.uyap_ref(mahkeme, no)
    return f"excel:dosya:{_satir_kimligi(satir)}"


def doldurma_onerisi(satir: dict, d: dict, dosyalar: list[dict]) -> dict | None:
    """isbu'da BOŞ olan esas no / mahkeme / özet alanlarını Excel'den doldurur (dolu alana dokunmaz)."""
    p = {}
    if not d.get("caseNumber") and _s(satir, "Dosya No"):
        p["case_number"] = _s(satir, "Dosya No")[:100]
    if not d.get("courtName") and _s(satir, "Mahkeme / Kurum"):
        p["court_name"] = _s(satir, "Mahkeme / Kurum")[:255]
    if not d.get("description") and _s(satir, "Özet"):
        p["description"] = sade.ozet_metni(_s(satir, "Özet"))
    if not p:
        return None
    ref = d.get("externalRef") or _dis_kimlik(satir)
    if ref.startswith("excel:") and set(p) == {"description"}:
        return None  # yalnız özet için davaya Excel kimliği bağlanmaz (UYAP bağı sonra gelsin)
    return ie.oge("case_upsert", ref, d["id"], p, f"{_ozet(satir)} · eksik: {', '.join(ETIKET[a] for a in p)}",
                  _gerekce(satir), "asistan", TUZ)


def siradaki_onerisi(satir: dict, d: dict) -> dict | None:
    """Excel 'Sıradaki İş' → dosyanın 'Beklenen' (waiting_for) alanı; aşama Excel 'Aşama'. Avukatın yazdığı
    beklenen varsa dokunulmaz."""
    is_ = _s(satir, "Sıradaki İş")
    if not is_ or ((d.get("workspace") or {}).get("waitingFor") or "").strip():
        return None
    p = {"target": "workspace", "waiting_for": sade.tek_satir(is_, 500)}
    if _s(satir, "Aşama"):
        p["stage"] = _s(satir, "Aşama")[:120]
    return ie.oge("next_step", f"excel:siradaki:{_satir_kimligi(satir)}", d["id"], p,
                  f"Sıradaki iş: {sade.tek_satir(is_, 200)} · {_s(satir, 'Müvekkil')}",
                  f"Ortak Excel satır {satir['_satir']}.", "asistan", TUZ)


def yeni_dava_onerisi(satir: dict, dosyalar: list[dict]) -> dict:
    tur, ozel = sade.dava_turu(_s(satir, "Konu"))
    p = {"title": _baslik(satir, dosyalar), "client_name": _s(satir, "Müvekkil")[:255], "case_type": tur}
    if ozel:
        p["custom_case_type"] = ozel
    if _s(satir, "Mahkeme / Kurum"):
        p["court_name"] = _s(satir, "Mahkeme / Kurum")[:255]
    if _s(satir, "Dosya No"):
        p["case_number"] = _s(satir, "Dosya No")[:100]
    if _s(satir, "Özet"):
        p["description"] = sade.ozet_metni(_s(satir, "Özet"))
    ref = _dis_kimlik(satir)
    return ie.oge("case_upsert", ref, {"external_ref": ref}, p, f"Yeni dava · {_ozet(satir)}", _gerekce(satir),
                  "asistan", TUZ)


def is_onerisi(is_satiri: dict, d: dict | None) -> dict | None:
    baslik = _s(is_satiri, "Yapılacak İş")
    if not baslik or _s(is_satiri, "Durum") == "Tamamlandı":
        return None
    p = {"title": sade.tek_satir(baslik, 500), "category": "cmk" if "CMK" in _s(is_satiri, "Müvekkil") else "dava"}
    if re.match(r"^\d{4}-\d{2}-\d{2}$", _s(is_satiri, "Son Gün")):
        p["due_date"] = _s(is_satiri, "Son Gün")
    if _s(is_satiri, "Not"):
        p["description"] = sade.tek_satir(_s(is_satiri, "Not"), 5000)
    ozet = f"İş: {sade.tek_satir(baslik, 200)} · {_s(is_satiri, 'Müvekkil')}"
    if p.get("due_date"):
        ozet += f" · son gün {date.fromisoformat(p['due_date']):%d.%m.%Y}"
    sorumlu = _s(is_satiri, "Sorumlu")
    gerekce = f"Ortak Excel İŞLER satır {is_satiri['_satir']}" + (f"; sorumlu {sorumlu}." if sorumlu else ".")
    ref = f"excel:is:{_kisa_hash(_s(is_satiri, 'Müvekkil'), baslik)}"  # satır numarasına bağlı değil
    return ie.oge("task_add", ref, d["id"] if d else None, p, ozet, gerekce, "asistan", TUZ)


def bekleyen_sadelestir(p: dict, adaylar: list[dict], dosyalar: list[dict], dava: dict | None = None) -> dict:
    """Kuyruktaki eski öneriyi Excel ölçüsüne indirir; aynı dış kimlikle gönderilince bekleyen öneri güncellenir.
    Var olan davaya yönelik öneride payload'a dokunulmaz, yalnız başlık/gerekçe sadeleşir."""
    pl = dict(p.get("payload") or {})
    ref = p.get("externalRef") or (p.get("caseRef") or {}).get("external_ref") or ""
    if p.get("caseId"):
        alanlar = ", ".join(ETIKET.get(a, a) for a in pl if a != "overwrite") or "UYAP dosyasıyla bağlantı"
        if len(adaylar) == 1:
            ozet, gerekce = f"{_ozet(adaylar[0])} · eksik: {alanlar}", _gerekce(adaylar[0])
        else:
            ad = (dava or {}).get("clientName") or (dava or {}).get("title") or p.get("caseTitle") or ""
            ozet = f"{sade.ozet_satiri(ad, None, pl.get('court_name'), pl.get('case_number'))} · eksik: {alanlar}"
            gerekce = "Büro kayıtlarındaki eksik alanlar."
        return ie.oge(p["kind"], ref, p["caseId"], pl, ozet, gerekce, p.get("source") or "uyap", TUZ)
    case_ref = p.get("caseRef") or {"external_ref": ref}
    if len(adaylar) == 1:
        s = adaylar[0]
        pl["title"] = _baslik(s, dosyalar)
        if _s(s, "Özet"):
            pl["description"] = sade.ozet_metni(_s(s, "Özet"))
        tur, ozel = sade.dava_turu(_s(s, "Konu"))
        pl["case_type"] = tur
        if ozel:
            pl["custom_case_type"] = ozel
        if not pl.get("court_name") and _s(s, "Mahkeme / Kurum"):
            pl["court_name"] = _s(s, "Mahkeme / Kurum")
        if not pl.get("case_number") and _s(s, "Dosya No"):
            pl["case_number"] = _s(s, "Dosya No")
        ozet, gerekce = f"Yeni dava · {_ozet(s)}", _gerekce(s)
    elif adaylar:
        ozet = f"Yeni dava · {pl.get('title', '')} · Excel'de {len(adaylar)} benzer satır var (satır " \
               f"{', '.join(str(s['_satir']) for s in adaylar)})"
        gerekce = "Hangi Excel satırı olduğu belirsiz; onaylamadan önce bakın."
    else:
        # Excel yalnız Mücahit'le ortak dosyaları tutar (avukat 22.10.2026); Excel'de olmamak ret sebebi DEĞİL.
        yer = " ".join(x for x in (pl.get("court_name"), pl.get("case_number")) if x)
        ozet = f"Yeni dava (UYAP) · {pl.get('title', '')}" + (f" · {yer}" if yer else "")
        gerekce = "UYAP'ta açık dosya; isbu'da kaydı yok. Takip edecekseniz onaylayın."
    if not pl.get("description"):
        pl.pop("description", None)
    return ie.oge(p["kind"], ref, case_ref, pl, ozet, gerekce, p.get("source") or "uyap", TUZ)


def eski_anahtar_isareti(p: dict) -> dict:
    """Önceki çalıştırmadan kalan, artık üretilmeyen Excel anahtarlı bekleyen öneri: aynı anahtarla 'reddedin' yazılır."""
    pl = dict(p.get("payload") or {})
    ref = p.get("externalRef") or ""
    case_ref = p["caseId"] if p.get("caseId") else (p.get("caseRef") or None)
    return ie.oge(p["kind"], ref, case_ref, pl, f"Eski anahtar · reddedin: {sade.tek_satir(p.get('summary'), 150)}",
                  "Aynı iş yeni anahtarla yeniden önerildi; bu kaydı reddedin.", p.get("source") or "asistan", TUZ)


# --- plan --------------------------------------------------------------------------------------------------------

class Plan:
    def __init__(self) -> None:
        self.items: list[dict] = []
        self.idx: dict[tuple, int] = {}      # (tür, dış kimlik) -> items konumu
        self.bek_idx: dict[tuple, int] = {}  # (tür, dış kimlik) -> bekleyen rapor konumu
        self.bek: list[dict] = []
        self.hatalar: list[str] = []

    def koy(self, it: dict) -> int:
        k = (it["kind"], it["external_ref"])
        if k in self.idx:
            self.items[self.idx[k]] = it
        else:
            self.idx[k] = len(self.items)
            self.items.append(it)
        return self.idx[k]

    def ekle(self, it: dict | None, r: dict) -> None:
        """Öneriyi plana koyar; aynı (tür, dış kimlik) kuyrukta bekliyorsa onunla BİRLEŞTİRİR (payload alanları
        birleşir; Excel satırı var olan davaya bağlıysa dava kimliği Excel'inki olur, yeni dava açılmaz)."""
        if not it:
            return
        k = (it["kind"], it["external_ref"])
        i = self.idx.get(k)
        if i is None:
            self.koy(it)
            r["oneriler"].append(it["summary"])
            return
        eski = self.items[i]
        case_ref = it.get("case_ref") if (it.get("case_ref") or {}).get("id") else eski.get("case_ref")
        birlesik = ie.oge(it["kind"], it["external_ref"], case_ref, {**eski["payload"], **it["payload"]},
                          it["summary"], it["reason"], eski["source"], TUZ)
        self.items[i] = birlesik
        if k in self.bek_idx:
            self.bek[self.bek_idx[k]]["yeni"] = birlesik["summary"]
        r["oneriler"].append(f"(kuyruktaki öneriyle birleşti) {birlesik['summary']}")


def plan_kur(excel: dict, davalar: list[dict], bekleyenler: list[dict]) -> dict:
    dosyalar = [s for s in excel["dosyalar"] if _s(s, "Durum") != "Tamamlandı" and _s(s, "Müvekkil")
                and "teyit edilecek" not in _s(s, "Müvekkil").lower()]
    P, kayit, belirsiz, yeni_yok, isler = Plan(), [], [], [], []
    dava_idx = {d["id"]: d for d in davalar}
    # 1) kuyruktaki eski öneriler önce sadeleşir; aynı (tür, dış kimlik) için Excel'den ikinci öğe üretilmez
    eski_excel_anahtarli = []
    for p in bekleyenler:
        if (p.get("externalRef") or "").startswith("excel:"):
            eski_excel_anahtarli.append(p)
            continue
        if p.get("kind") != "case_upsert":
            continue
        d = dava_idx.get(p.get("caseId") or "")
        aday = bekleyen_bul(p, dosyalar, d)
        try:
            it = bekleyen_sadelestir(p, aday, dosyalar, d)
        except ValueError as h:
            P.hatalar.append(f"bekleyen {p['id'][:8]}: {h}")
            continue
        P.bek_idx[(it["kind"], it["external_ref"])] = len(P.bek)
        P.koy(it)
        P.bek.append({"id": p["id"][:8], "eski": p.get("summary"), "yeni": it["summary"], "excel": [s["_satir"] for s in aday]})
    # yeni dava freni: kuyrukta aynı müvekkil için UYAP kaynaklı, Excel satırıyla eşleşmemiş 'Yeni dava' varsa
    eslesmis_bek = {b["id"] for b in P.bek if len(b["excel"]) == 1}
    bekleyen_muvekkil = [(p.get("payload") or {}).get("client_name") or "" for p in bekleyenler
                         if p.get("kind") == "case_upsert" and not p.get("caseId")
                         and not (p.get("externalRef") or "").startswith("excel:") and p["id"][:8] not in eslesmis_bek]
    # 2) Excel DOSYALAR
    kimlikler: dict[str, list[int]] = {}
    for s in dosyalar:
        kimlikler.setdefault(_satir_kimligi(s), []).append(s["_satir"])
    eslesen_dava: dict[str, int] = {}
    for s in dosyalar:
        aday, yol = dava_bul(s, davalar)
        r = {"satir": s["_satir"], "muvekkil": _s(s, "Müvekkil"), "mahkeme": _s(s, "Mahkeme / Kurum"),
             "no": _s(s, "Dosya No"), "yol": yol, "isbu": [d["title"] for d in aday], "oneriler": []}
        ayni_muvekkil = [x for x in dosyalar if _kelimeler(_s(x, "Müvekkil")) == _kelimeler(_s(s, "Müvekkil"))]
        try:
            if len(kimlikler[_satir_kimligi(s)]) > 1:
                r["durum"], r["neden"] = "belirsiz", f"aynı anahtar (satır {kimlikler[_satir_kimligi(s)]}); Dosya No veya Konu ayırt etmiyor"
                belirsiz.append(r)
            elif len(aday) > 1:
                r["durum"] = "belirsiz"
                belirsiz.append(r)
            elif aday and yol == "ad" and len(ayni_muvekkil) > 1:
                r["durum"], r["neden"] = "belirsiz", f"müvekkilin Excel'de {len(ayni_muvekkil)} dosyası var; isbu davası yalnız adla bulundu"
                belirsiz.append(r)
            elif aday and aday[0]["id"] in eslesen_dava:
                r["durum"], r["neden"] = "belirsiz", f"isbu davası satır {eslesen_dava[aday[0]['id']]} ile zaten eşleşti"
                belirsiz.append(r)
            elif aday:
                d = aday[0]
                eslesen_dava[d["id"]] = s["_satir"]
                r["durum"] = "eşleşti"
                P.ekle(doldurma_onerisi(s, d, dosyalar), r)
                P.ekle(siradaki_onerisi(s, d), r)
            else:
                benzer = [d["title"] for d in davalar if ad_benzer(_s(s, "Müvekkil"), d.get("clientName"), d.get("title"))]
                kuyruk = [m for m in bekleyen_muvekkil if ad_eslesir(_s(s, "Müvekkil"), m)]
                if benzer or kuyruk:
                    r["durum"] = "isbu'da benzer ad var, yeni dava önerilmedi" if benzer else \
                        "kuyrukta aynı müvekkilin UYAP önerisi var, yeni dava önerilmedi"
                    r["benzer"] = (benzer or kuyruk)[:5]
                    yeni_yok.append(r)
                else:
                    r["durum"] = "yeni dava"
                    P.ekle(yeni_dava_onerisi(s, dosyalar), r)
        except ValueError as h:
            r["durum"], r["hata"] = "hata", str(h)
            P.hatalar.append(f"satır {s['_satir']}: {h}")
        kayit.append(r)
    # 3) İŞLER
    for i in excel["isler"]:
        if _s(i, "Durum") == "Tamamlandı" or not _s(i, "Yapılacak İş"):
            continue
        aday = [d for d in davalar if ad_eslesir(_s(i, "Müvekkil"), d.get("clientName"), d.get("title"))]
        try:
            it = is_onerisi(i, aday[0] if len(aday) == 1 else None)
        except ValueError as h:
            P.hatalar.append(f"İŞLER satır {i['_satir']}: {h}")
            continue
        if it:
            P.koy(it)
            isler.append({"satir": i["_satir"], "muvekkil": _s(i, "Müvekkil"), "dava": aday[0]["title"] if len(aday) == 1
                          else ("belirsiz" if aday else "davasız"), "ozet": it["summary"]})
    # 4) önceki çalıştırmanın artık üretilmeyen Excel anahtarlı bekleyenleri
    eski = []
    for p in eski_excel_anahtarli:
        if (p["kind"], p["externalRef"]) in P.idx:
            continue
        try:
            it = eski_anahtar_isareti(p)
        except ValueError as h:
            P.hatalar.append(f"eski anahtar {p['id'][:8]}: {h}")
            continue
        P.koy(it)
        eski.append({"id": p["id"][:8], "eski": p.get("summary")})
    return {"items": P.items, "dosyalar": kayit, "belirsiz": belirsiz, "yeni_yok": yeni_yok, "isler": isler,
            "bekleyen": P.bek, "eski_anahtar": eski, "hatalar": P.hatalar}


def _sayim(items: list[dict]) -> dict:
    out = {}
    for it in items:
        out[it["kind"]] = out.get(it["kind"], 0) + 1
    return out


def md_plan(plan: dict, bugun: date, gonderildi: bool) -> str:
    s = [f"# Excel → isbu-ofis sade öneri planı · {bugun:%d.%m.%Y}", ""]
    s.append("**GÖNDERİLDİ.**" if gonderildi else "**KURU ÇALIŞMA — gönderilmedi.** Gönderim `plan --gonder`.")
    s.append("")
    s.append(f"- Öneri: **{len(plan['items'])}** · tür bazında: "
             + ", ".join(f"{k} {v}" for k, v in sorted(_sayim(plan["items"]).items())))
    d = plan["dosyalar"]
    s.append(f"- Excel açık dosya: {len(d)} · isbu'da bulunan: {sum(1 for r in d if r['durum'] == 'eşleşti')} · "
             f"belirsiz: {len(plan['belirsiz'])} · yeni dava: {sum(1 for r in d if r['durum'] == 'yeni dava')} · "
             f"benzer ad/kuyruk nedeniyle önerilmeyen: {len(plan['yeni_yok'])} · hata: {len(plan['hatalar'])}")
    s.append(f"- İŞLER sayfasından görev: {len(plan['isler'])} · kuyruktaki eski öneri sadeleştirilen: "
             f"{len(plan['bekleyen'])} (Excel'de olan {sum(1 for b in plan['bekleyen'] if len(b['excel']) == 1)}) · "
             f"eski anahtarlı 'reddedin' işareti: {len(plan['eski_anahtar'])}")
    if plan["hatalar"]:
        s += ["", "## Hatalar (öneri üretilmedi)", ""] + [f"- {h}" for h in plan["hatalar"]]
    s += ["", "## Excel dosyaları", "", "| Satır | Müvekkil | Mahkeme / No | Eşleşme | isbu davası | Öneri |",
          "|---|---|---|---|---|---|"]
    for r in d:
        isbu = "; ".join(r["isbu"]) or ("benzer: " + ", ".join(r["benzer"]) if r.get("benzer") else "")
        s.append(f"| {r['satir']} | {r['muvekkil']} | {r['mahkeme']} {r['no']} | {r['durum']} ({r['yol']})"
                 f"{' — ' + r['neden'] if r.get('neden') else ''}{' — ' + r['hata'] if r.get('hata') else ''} | "
                 f"{isbu} | {'<br>'.join(r['oneriler']) or '—'} |")
    s += ["", "## İŞLER", "", "| Satır | Müvekkil | Dava | Öneri |", "|---|---|---|---|"]
    s += [f"| {i['satir']} | {i['muvekkil']} | {i['dava']} | {i['ozet']} |" for i in plan["isler"]]
    s += ["", "## Kuyruktaki eski öneriler (sadeleştirme)", "", "| Öneri | Eski başlık | Yeni başlık | Excel satır |",
          "|---|---|---|---|"]
    s += [f"| {b['id']} | {b['eski']} | {b['yeni']} | {', '.join(map(str, b['excel'])) or '—'} |" for b in plan["bekleyen"]]
    if plan["eski_anahtar"]:
        s += ["", "## Eski anahtarlı bekleyenler ('reddedin' işareti)", ""] + [f"- {e['id']} · {e['eski']}" for e in plan["eski_anahtar"]]
    return "\n".join(s) + "\n"


def _canli_excel_cek() -> None:
    """Canlı Excel OneDrive'da (22.10.2026). Paylaşım linki (`buro_excel_onedrive.paylasim`) yalnız gitignore'daki
    config/paths.local.json içinde tanımlıysa her çalıştırmada önce canlı nüsha indirilir; indirilemezse eldeki
    kopyayla sürer, uyarı yazar."""
    from paths import _load_cfg
    link = (_load_cfg().get("buro_excel_onedrive") or {}).get("paylasim")
    if not link:
        print("UYARI: OneDrive paylaşım linki tanımsız (config buro_excel_onedrive.paylasim); eldeki kopya kullanılıyor.")
        return
    try:
        import excel_indir
        from paths import buro_excel as _be
        excel_indir.yerlestir(excel_indir.paylasim_indir(link), _be(), kuru=False)
    except (RuntimeError, OSError) as h:
        print(f"UYARI: canlı Excel indirilemedi ({h}); eldeki kopya kullanılıyor.")


def calistir(cikti: Path, gonder: bool, excel_yolu: Path | None = None, api=None, bugun: date | None = None) -> dict:
    import isbu_ofis
    from paths import buro_excel
    bugun = bugun or date.today()
    if not excel_yolu:
        _canli_excel_cek()
    excel = excel_oku(excel_yolu or buro_excel())
    api = api or isbu_ofis.IsbuOfis()
    davalar = api.cases(include_closed=False)
    bekleyenler = api.proposals("pending")
    plan = plan_kur(excel, davalar, bekleyenler)
    sonuc = None
    if gonder and plan["items"]:
        sonuc = api.post_proposals(plan["items"], gonder=True)
    ie._yaz(cikti, "plan.json", {"bugun": bugun.isoformat(), "gonderildi": bool(sonuc), "sonuc": sonuc, **plan})
    ie._yaz(cikti, "plan.md", md_plan(plan, bugun, bool(sonuc)))
    return {"plan": plan, "sonuc": sonuc}


def main(argv: list[str]) -> int:
    ap = argparse.ArgumentParser(prog="isbu_excel.py")
    alt = ap.add_subparsers(dest="komut", required=True)
    p = alt.add_parser("plan")
    p.add_argument("--gonder", action="store_true", help="önerileri isbu-ofis kuyruğuna yazar (avukat uygulamada onaylar)")
    p.add_argument("--cikti")
    p.add_argument("--excel")
    a = ap.parse_args(argv)
    cikti = Path(a.cikti) if a.cikti else REPO_ROOT / "tmp" / "isbu-ofis" / f"{date.today().isoformat()}-excel"
    r = calistir(cikti, a.gonder, Path(a.excel) if a.excel else None)
    plan = r["plan"]
    print(f"Öneri: {len(plan['items'])} {_sayim(plan['items'])} · belirsiz {len(plan['belirsiz'])} · "
          f"benzer-ad freni {len(plan['yeni_yok'])} · görev {len(plan['isler'])} · eski öneri {len(plan['bekleyen'])} · "
          f"eski anahtar {len(plan['eski_anahtar'])} · hata {len(plan['hatalar'])}")
    if r["sonuc"]:
        for s in r["sonuc"]:
            print("gönderim:", s.get("http"), s.get("counts"))
    print("rapor:", cikti / "plan.md")
    return 0


if __name__ == "__main__":
    sys.exit(main(sys.argv[1:]))
