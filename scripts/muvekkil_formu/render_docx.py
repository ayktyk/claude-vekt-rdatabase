"""DOCX üretici — Vega Hukuk müvekkil bilgi formu (telefonda/Word'de yazılabilir).

Cevap hücreleri boş bırakılır; müvekkil hücreye yazar. Seçenekler ☐ karakteriyle
verilir; müvekkil kutuyu ☒ yapar veya yanına X yazar.
"""
from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

from . import brand, schema

BOX = "☐"  # ☐


def _rgb(hexs: str) -> RGBColor:
    return RGBColor.from_string(hexs.lstrip("#").upper())


NAVY, GOLD, CREAM, INK, MUTED, LINE, GREY = (
    _rgb(brand.NAVY), _rgb(brand.GOLD), _rgb(brand.CREAM), _rgb(brand.INK),
    _rgb(brand.MUTED), _rgb(brand.LINE), _rgb(brand.GREY_BG),
)


# ---- OXML yardımcıları -------------------------------------------------------
def shade(cell, hexs: str) -> None:
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hexs.lstrip("#"))
    tcPr.append(shd)


def borders(table, hexs: str = brand.LINE, sz: int = 4, inside: bool = True) -> None:
    tblPr = table._tbl.tblPr
    b = OxmlElement("w:tblBorders")
    edges = ["top", "left", "bottom", "right"] + (["insideH", "insideV"] if inside else [])
    for e in edges:
        el = OxmlElement(f"w:{e}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), str(sz))
        el.set(qn("w:space"), "0")
        el.set(qn("w:color"), hexs.lstrip("#"))
        b.append(el)
    tblPr.append(b)


def no_borders(table) -> None:
    tblPr = table._tbl.tblPr
    b = OxmlElement("w:tblBorders")
    for e in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{e}")
        el.set(qn("w:val"), "nil")
        b.append(el)
    tblPr.append(b)


def cell_margins(table, top=60, bottom=60, left=90, right=90) -> None:
    tblPr = table._tbl.tblPr
    m = OxmlElement("w:tblCellMar")
    for k, v in (("top", top), ("left", left), ("bottom", bottom), ("right", right)):
        el = OxmlElement(f"w:{k}")
        el.set(qn("w:w"), str(v))
        el.set(qn("w:type"), "dxa")
        m.append(el)
    tblPr.append(m)


def row_height(row, cm: float) -> None:
    trPr = row._tr.get_or_add_trPr()
    h = OxmlElement("w:trHeight")
    h.set(qn("w:val"), str(int(cm * 567)))
    h.set(qn("w:hRule"), "atLeast")
    trPr.append(h)


def para_border_bottom(p, hexs: str, sz: int = 8) -> None:
    pPr = p._p.get_or_add_pPr()
    pb = OxmlElement("w:pBdr")
    el = OxmlElement("w:bottom")
    el.set(qn("w:val"), "single")
    el.set(qn("w:sz"), str(sz))
    el.set(qn("w:space"), "2")
    el.set(qn("w:color"), hexs.lstrip("#"))
    pb.append(el)
    pPr.append(pb)


def field_code(run, code: str) -> None:
    for tag, txt in (("begin", None), (None, code), ("separate", None), ("end", None)):
        if tag:
            el = OxmlElement("w:fldChar")
            el.set(qn("w:fldCharType"), tag)
        else:
            el = OxmlElement("w:instrText")
            el.set(qn("xml:space"), "preserve")
            el.text = txt
        run._r.append(el)


# ---- yazı yardımcıları --------------------------------------------------------
def run(p, text, size=9.5, bold=False, color=INK, font="Segoe UI", italic=False):
    r = p.add_run(text)
    r.font.name = font
    r._element.rPr.rFonts.set(qn("w:eastAsia"), font)
    r.font.size = Pt(size)
    r.bold = bold
    r.italic = italic
    r.font.color.rgb = color
    return r


def tight(p, before=0, after=0, line=1.05):
    pf = p.paragraph_format
    pf.space_before = Pt(before)
    pf.space_after = Pt(after)
    pf.line_spacing = line
    return p


def label_cell(cell, label, hint=None, fill=brand.CREAM):
    shade(cell, fill)
    p = tight(cell.paragraphs[0])
    run(p, label, 9, bold=True, color=INK)
    if hint:
        run(tight(cell.add_paragraph()), hint, 7.5, color=MUTED)


def answer_cell(cell, lines=1):
    tight(cell.paragraphs[0])
    for _ in range(lines - 1):
        tight(cell.add_paragraph())


def boxes(p, options, size=9.5):
    for i, o in enumerate(options):
        run(p, BOX + " ", size + 1, color=NAVY, font="Segoe UI Symbol")
        run(p, o.replace(" ", " "), size)
        if i < len(options) - 1:
            run(p, "      ", size)


# ---- belge -----------------------------------------------------------------------
class DocxForm:
    def __init__(self):
        self.doc = Document()
        self.buro = brand.buro()
        sec = self.doc.sections[0]
        sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
        sec.left_margin = sec.right_margin = Cm(1.8)
        sec.top_margin, sec.bottom_margin = Cm(1.6), Cm(1.6)
        sec.header_distance = sec.footer_distance = Cm(0.7)
        st = self.doc.styles["Normal"]
        st.font.name = "Segoe UI"
        st.element.rPr.rFonts.set(qn("w:eastAsia"), "Segoe UI")
        st.font.size = Pt(9.5)
        st.paragraph_format.space_after = Pt(0)
        self.width = sec.page_width - sec.left_margin - sec.right_margin
        self._header_footer(sec)

    # -- iskelet
    def _header_footer(self, sec) -> None:
        hp = sec.header.paragraphs[0]
        t = sec.header.add_table(1, 1, self.width)
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        no_borders(t)
        cell_margins(t, 80, 80, 140, 140)
        c = t.rows[0].cells[0]
        c.width = self.width
        shade(c, brand.NAVY)
        p1 = tight(c.paragraphs[0])
        run(p1, brand.WORDMARK, 14, bold=True, color=CREAM, font="Georgia")
        run(p1, "	", 8)
        run(p1, schema.META["baslik"], 8.5, bold=True, color=CREAM)
        p2 = tight(c.add_paragraph())
        run(p2, "  ".join(brand.TAGLINE), 5.5, color=GOLD)
        run(p2, "	", 7)
        run(p2, schema.META["alt_baslik"], 7.5, color=CREAM)
        for p in (p1, p2):
            p.paragraph_format.tab_stops.add_tab_stop(self.width - Cm(0.5), WD_ALIGN_PARAGRAPH.RIGHT)
        hp._p.getparent().remove(hp._p)

        fp = sec.footer.paragraphs[0]
        tight(fp)
        para_border_bottom(fp, brand.GOLD, 6)
        b = self.buro
        f2 = tight(sec.footer.add_paragraph(), before=2)
        run(f2, f'{b["avukat"]}  ·  {b["buro"]}  ·  {b["adres"]}', 6.8, color=MUTED)
        run(f2, "\t", 6.8)
        run(f2, "Sayfa ", 7, bold=True, color=NAVY)
        field_code(run(f2, "", 7, bold=True, color=NAVY), "PAGE")
        run(f2, " / ", 7, bold=True, color=NAVY)
        field_code(run(f2, "", 7, bold=True, color=NAVY), "NUMPAGES")
        f3 = tight(sec.footer.add_paragraph())
        run(f3, f'{b["site"]}  ·  {b["eposta"]}  ·  {b["baro"]}', 6.5, color=MUTED)
        run(f3, "\t", 6.5)
        run(f3, f'Form sürümü {schema.META["surum"]}', 6.5, color=MUTED)
        for p in (f2, f3):
            p.paragraph_format.tab_stops.add_tab_stop(self.width, WD_ALIGN_PARAGRAPH.RIGHT)

    def one_cell(self, fill=None, border=None, sz=6, pad=(100, 100, 160, 160)):
        t = self.doc.add_table(1, 1)
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        t.autofit = False
        t.columns[0].width = self.width
        t.rows[0].cells[0].width = self.width
        if border:
            borders(t, border, sz, inside=False)
        else:
            no_borders(t)
        cell_margins(t, *pad)
        c = t.rows[0].cells[0]
        if fill:
            shade(c, fill)
        return c

    def spacer(self, pt=4):
        p = tight(self.doc.add_paragraph(), after=pt, line=1.0)
        run(p, "", 1)
        p.paragraph_format.line_spacing = Pt(1)

    # -- kapak
    def cover(self) -> None:
        c = self.one_cell(fill=brand.NAVY, pad=(220, 220, 320, 320))
        run(tight(c.paragraphs[0]), schema.META["baslik"], 20, bold=True, color=CREAM, font="Georgia")
        run(tight(c.add_paragraph(), before=2), schema.META["alt_baslik"], 13, color=GOLD, font="Georgia")
        p = tight(c.add_paragraph(), before=4)
        p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        run(p, f'Sürüm {schema.META["surum"]}', 7, color=GOLD)
        self.spacer(6)
        run(tight(self.doc.add_paragraph(), after=6), schema.META["aciklama"], 9.5)

        c = self.one_cell(fill=brand.CREAM)
        run(tight(c.paragraphs[0], after=3), "FORMU DOLDURMADAN ÖNCE", 8.5, bold=True, color=NAVY)
        for i, t in enumerate(schema.TALIMATLAR, 1):
            p = tight(c.add_paragraph(), after=2)
            run(p, f"{i}.  ", 9, bold=True, color=NAVY)
            run(p, t, 9)
        p = tight(c.add_paragraph(), before=3)
        run(p, "Seçenekli sorularda uygun kutuyu ", 8, color=MUTED)
        run(p, "☒", 9, color=NAVY, font="Segoe UI Symbol")
        run(p, " yapın veya yanına X yazın. Cevapları boş hücrelere yazın.", 8, color=MUTED)
        self.spacer(4)

        su = schema.SURE_UYARISI
        c = self.one_cell(fill="#FFFFFF", border=brand.GOLD, sz=12)
        run(tight(c.paragraphs[0], after=3), su["baslik"], 9, bold=True, color=NAVY)
        for t in su["maddeler"]:
            p = tight(c.add_paragraph(), after=2)
            p.paragraph_format.left_indent = Cm(0.4)
            p.paragraph_format.first_line_indent = Cm(-0.4)
            run(p, "•  ", 9, color=GOLD)
            run(p, t, 9)
        run(tight(c.add_paragraph(), before=4), su["son_satir"], 9, bold=True, color=NAVY)
        self.spacer(6)

    # -- bölüm başlığı
    def section(self, kod, baslik) -> None:
        p = tight(self.doc.add_paragraph(), before=10, after=6)
        p.paragraph_format.keep_with_next = True
        r = run(p, f" {kod} ", 10, bold=True, color=NAVY, font="Georgia")
        rPr = r._element.rPr
        shd = OxmlElement("w:shd")
        shd.set(qn("w:val"), "clear")
        shd.set(qn("w:fill"), brand.GOLD.lstrip("#"))
        rPr.append(shd)
        run(p, "   " + baslik, 12.5, bold=True, color=NAVY, font="Georgia")
        para_border_bottom(p, brand.GOLD, 8)

    # -- alanlar
    def _grid(self, rows, cols, widths):
        t = self.doc.add_table(rows, cols)
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        t.autofit = False
        borders(t)
        cell_margins(t)
        for r in t.rows:
            for c, w in zip(r.cells, widths):
                c.width = w
        return t

    def f_text(self, label, hint=None) -> None:
        w1 = Cm(6.6)
        t = self._grid(1, 2, [w1, self.width - w1])
        label_cell(t.cell(0, 0), label, hint)
        answer_cell(t.cell(0, 1))
        self.spacer(3)

    def f_pair(self, a, b) -> None:
        w = (self.width - Cm(0.3)) / 2
        t = self._grid(2, 2, [w, w])
        for i, (lab, hint) in enumerate((a, b)):
            label_cell(t.cell(0, i), lab, hint)
            answer_cell(t.cell(1, i))
        row_height(t.rows[1], 0.75)
        self.spacer(3)

    def f_multi(self, label, lines, hint=None) -> None:
        t = self._grid(2, 1, [self.width])
        label_cell(t.cell(0, 0), label, hint)
        answer_cell(t.cell(1, 0), lines)
        row_height(t.rows[1], 0.55 * lines + 0.2)
        self.spacer(3)

    def f_yesno(self, label, hint=None) -> None:
        w1, w2 = Cm(6.6), Cm(3.2)
        t = self._grid(1, 3, [w1, w2, self.width - w1 - w2])
        label_cell(t.cell(0, 0), label, hint)
        boxes(tight(t.cell(0, 1).paragraphs[0]), ["Evet", "Hayır"])
        answer_cell(t.cell(0, 2))
        self.spacer(3)

    def f_choice(self, label, options, other=False) -> None:
        p = tight(self.doc.add_paragraph(), after=2)
        p.paragraph_format.keep_with_next = True
        run(p, label, 9, bold=True)
        p = tight(self.doc.add_paragraph(), after=2)
        p.paragraph_format.left_indent = Cm(0.2)
        boxes(p, options)
        if other:
            run(p, "     ", 9.5)
            run(p, BOX + " ", 10.5, color=NAVY, font="Segoe UI Symbol")
            run(p, "Diğer: ______________________", 9.5)
        self.spacer(4)

    def f_table(self, label, cols, nrows) -> None:
        n = len(cols)
        weights = [1.2, 1, 0.9, 1, 1.5] if n == 5 else ([0.8, 1, 1.6] if n == 3 else [1] * n)
        tot = sum(weights)
        widths = [int(self.width * w / tot) for w in weights]
        p = tight(self.doc.add_paragraph(), after=3)
        p.paragraph_format.keep_with_next = True
        run(p, label, 9, bold=True)
        t = self._grid(nrows + 1, n, widths)
        for i, col in enumerate(cols):
            c = t.cell(0, i)
            shade(c, brand.NAVY)
            run(tight(c.paragraphs[0]), col, 8, bold=True, color=CREAM)
        for r in range(1, nrows + 1):
            row_height(t.rows[r], 0.65)
            for i in range(n):
                answer_cell(t.cell(r, i))
        self.spacer(4)

    def f_docs(self, items) -> None:
        w = [self.width - Cm(2.2) - Cm(3.0) - Cm(1.5), Cm(2.2), Cm(3.0), Cm(1.5)]
        t = self._grid(len(items) + 1, 4, w)
        for i, col in enumerate(("Belge", "Elimde var", "Temin edebilirim", "Yok")):
            c = t.cell(0, i)
            shade(c, brand.NAVY)
            run(tight(c.paragraphs[0]), col, 8, bold=True, color=CREAM)
        for r, it in enumerate(items, 1):
            if r % 2 == 1:
                for i in range(4):
                    shade(t.cell(r, i), brand.CREAM)
            run(tight(t.cell(r, 0).paragraphs[0]), f"{r}. {it}", 8.5)
            for i in (1, 2, 3):
                p = tight(t.cell(r, i).paragraphs[0])
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                run(p, BOX, 11, color=NAVY, font="Segoe UI Symbol")
        self.spacer(4)

    def f_note(self, s) -> None:
        c = self.one_cell(fill=brand.CREAM)
        run(tight(c.paragraphs[0]), s, 8.5)
        self.spacer(4)

    # -- kapanış
    def beyan(self) -> None:
        c = self.one_cell(fill="#FFFFFF", border=brand.NAVY, sz=8)
        run(tight(c.paragraphs[0], after=3), "BEYAN VE İMZA", 9, bold=True, color=NAVY)
        run(tight(c.add_paragraph(), after=6), schema.BEYAN, 9)
        t = c.add_table(2, 3)
        borders(t)
        cell_margins(t)
        w = (self.width - Cm(0.9)) / 3
        for i, lab in enumerate(("Tarih", "Ad Soyad", "İmza")):
            t.cell(0, i).width = t.cell(1, i).width = w
            label_cell(t.cell(0, i), lab)
            answer_cell(t.cell(1, i))
        row_height(t.rows[1], 1.2)
        self.spacer(6)

    def buro_kutusu(self) -> None:
        bk = schema.BURO_KUTUSU
        c = self.one_cell(fill=brand.GREY_BG, border=brand.MUTED, sz=4)
        run(tight(c.paragraphs[0], after=4), bk["baslik"], 8.5, bold=True)
        t = c.add_table(len(bk["satirlar"]) * 2, 2)
        borders(t, brand.LINE, 2)
        cell_margins(t, 40, 40, 80, 80)
        w = (self.width - Cm(0.9)) / 2
        for r, (a, b) in enumerate(bk["satirlar"]):
            for i, lab in enumerate((a, b)):
                t.cell(2 * r, i).width = t.cell(2 * r + 1, i).width = w
                shade(t.cell(2 * r, i), brand.GREY_BG)
                run(tight(t.cell(2 * r, i).paragraphs[0]), lab, 7.5)
                shade(t.cell(2 * r + 1, i), "#FFFFFF")
                answer_cell(t.cell(2 * r + 1, i))
        for k in bk["kontrol"]:
            p = tight(c.add_paragraph(), before=3)
            run(p, BOX + "  ", 10, color=NAVY, font="Segoe UI Symbol")
            run(p, k, 8)

    # -- ana akış
    def build(self, out: Path) -> None:
        self.cover()
        for sec in schema.SECTIONS:
            self.section(sec["kod"], sec["baslik"])
            for f in sec["alanlar"]:
                kind = f[0]
                if kind == "text":
                    self.f_text(f[1], f[2] if len(f) > 2 else None)
                elif kind == "date":
                    self.f_text(f[1], f[2] if len(f) > 2 else "GG.AA.YYYY")
                elif kind == "pair":
                    self.f_pair(f[1], f[2])
                elif kind == "multi":
                    self.f_multi(f[1], f[2], f[3] if len(f) > 3 else None)
                elif kind == "choice":
                    self.f_choice(f[1], f[2], f[3] if len(f) > 3 else False)
                elif kind == "yesno":
                    self.f_yesno(f[1], f[2] if len(f) > 2 else None)
                elif kind == "table":
                    self.f_table(f[1], f[2], f[3])
                elif kind == "docs":
                    self.f_docs(f[1])
                elif kind == "note":
                    self.f_note(f[1])
                else:
                    raise ValueError(f"bilinmeyen alan türü: {kind}")
        self.spacer(6)
        self.beyan()
        self.buro_kutusu()
        cp = self.doc.core_properties
        cp.title = f"{schema.META['baslik']} — {schema.META['alt_baslik']}"
        cp.author = self.buro["avukat"]
        out.parent.mkdir(parents=True, exist_ok=True)
        self.doc.save(str(out))


def render(out: Path) -> None:
    DocxForm().build(out)
