#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""isbu_esitle.py — büro kaynakları ↔ isbu-ofis eşleştirme ve öneri planı (Aşama 2, kuru çalışma).

Plan: docs/superpowers/specs/2026-10-10-isbu-ofis-entegrasyon-plani.md · Sözleşme: isbu-ofis docs/ENTEGRASYON.md.
UYAP'a ve Yargı PRO uygulamasına İSTEK ATMAZ: yalnız diskteki büro verisi okunur —
  UYAP önbelleği (00-uyap-onbellek.json), eşleşme tablosu (00-UYAP-ESLESME-TABLOSU.md), CMK listesi
  (00-CMK-DOSYALARI.md), Aktif Davalar altındaki klon manifestleri (uyap-project.json), dava klasörleri
  (YYYY-NNN), iş takip Excel'i (yalnız okuma), Süre Kalkanı depoları (00-Sure-Kalkani.json) ve icra_ozet.
isbu-ofis tarafında yalnız GET çağrılır. Öneri gönderimi (`plan --gonder`) avukat onayı olmadan çalıştırılmaz.

Eşleştirme kuralı (tahmin YOK): isbu davasında external_ref veya buro_kodu varsa önce onunla; yoksa esas no +
mahkeme normalize edilerek (uyap_ortak.ayni_dosya). Tek-tek karşılık yoksa (birden çok aday, mahkeme adı eksik,
klasör kodu çakışması) dosya "belirsiz" listesine gider ve öneri üretilmez.
Süre önerisi yalnız Süre Kalkanı satırından üretilir; kaynak yoksa öneri yok, raporda "kaynak yok" yazar.

Kullanım:
  python scripts/isbu_esitle.py eslestir [--cikti <klasör>]
  python scripts/isbu_esitle.py plan     [--cikti <klasör>] [--gonder]
Çıktı (varsayılan): tmp/isbu-ofis/<tarih>/ (gitignore'da; müvekkil verisi içerir).
"""
from __future__ import annotations

import argparse
import json
import os
import re
import sys
from datetime import date, datetime, timedelta, timezone
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))
import isbu_sema  # noqa: E402
import uyap_ortak as uo  # noqa: E402

REPO_ROOT = Path(__file__).resolve().parent.parent
TR_SAAT = timezone(timedelta(hours=3))  # Türkiye 2016'dan beri sabit UTC+3
KLASOR_RE = re.compile(r"^(\d{4}-\d{3})(?!\d)")
SURE_DIKKAT = ("Karar", "Tebligat", "Tensip", "Bilirkişi")
BUDANAN = {"evraklar", "node_modules", ".git"}


# --- yardımcılar ------------------------------------------------------------------------------------------------

def norm_mahkeme(ad) -> str:
    return re.sub(r"\s+", " ", re.sub(r"[^0-9a-z.() ]+", " ", uo.tr(ad))).strip()


def uyap_ref(mahkeme, no, portal: str = "avukat") -> str:
    """Davanın UYAP dış kimliği (normalize: büyük/küçük harf ve Türkçe karakter farkı kimliği değiştirmez)."""
    return f"uyap:{portal or 'avukat'}:{norm_mahkeme(mahkeme)}|{uo.dosya_no(no)}"[:200]


def klasor_kodu(ad) -> str:
    m = KLASOR_RE.match(str(ad or "").strip())
    return m.group(1) if m else ""


def tr_iso(metin) -> str | None:
    """UYAP tarih-saati -> '+03:00' dilimli ISO. Dilimsiz değer Türkiye yerel saati sayılır; tarih yoksa None."""
    s = str(metin or "").strip()
    if not s or len(s) < 16:
        return None  # yalnız gün: saat bilinmiyor, uydurulmaz
    try:
        t = datetime.fromisoformat(s.replace("Z", "+00:00"))
    except ValueError:
        return None
    t = t.astimezone(TR_SAAT) if t.tzinfo else t.replace(tzinfo=TR_SAAT)
    return t.replace(microsecond=0).isoformat()


def ayni_dakika(a, b) -> bool:
    try:
        ta = datetime.fromisoformat(str(a).replace("Z", "+00:00"))
        tb = datetime.fromisoformat(str(b).replace("Z", "+00:00"))
    except ValueError:
        return False
    if not (ta.tzinfo and tb.tzinfo):
        return False
    return abs((ta - tb).total_seconds()) < 60


# --- büro envanteri ---------------------------------------------------------------------------------------------

def dava_klasorleri(kok: Path) -> dict[str, list[str]]:
    """{kod: [klasör adı, ...]} — aynı koda birden çok klasör çakışma sayılır."""
    out: dict[str, list[str]] = {}
    if kok.is_dir():
        for p in sorted(kok.iterdir()):
            kod = klasor_kodu(p.name) if p.is_dir() else ""
            if kod:
                out.setdefault(kod, []).append(p.name)
    return out


def klon_tara(kok: Path) -> list[dict]:
    """Aktif Davalar altındaki uyap-project.json manifestleri (evrak klasörlerine inilmez)."""
    out = []
    for dizin, altlar, dosyalar in os.walk(kok):
        altlar[:] = [a for a in altlar if a not in BUDANAN and not a.endswith("_ekler")]
        if "uyap-project.json" not in dosyalar:
            continue
        yol = Path(dizin)
        try:
            proje = json.loads((yol / "uyap-project.json").read_text(encoding="utf-8"))
        except (ValueError, OSError):
            continue
        goreli = yol.relative_to(kok).parts
        out.append({"mahkeme": proje.get("mahkeme"), "no": uo.dosya_no(proje.get("esasNo")),
                    "portal": proje.get("portal") or "avukat", "icra": bool(proje.get("isIcra")), "klon": str(yol),
                    "klasor": klasor_kodu(goreli[0]) if goreli else ""})
        altlar[:] = []  # klonun içinde başka klon aranmaz
    return out


def excel_satirlari(yol: Path) -> tuple[list[dict], str]:
    """İş takip Excel'i DOSYALAR sayfası (yalnız okuma; openpyxl). (satırlar, uyarı)."""
    if not yol.is_file():
        return [], f"Excel yok: {yol.name}"
    try:
        import openpyxl
        wb = openpyxl.load_workbook(yol, read_only=True, data_only=True)
    except Exception as hata:  # noqa: BLE001 — Excel okunamazsa diğer kaynaklarla devam edilir
        return [], f"Excel okunamadı: {type(hata).__name__}"
    try:
        if "DOSYALAR" not in wb.sheetnames:
            return [], "Excel'de DOSYALAR sayfası yok"
        out = []
        for satir in wb["DOSYALAR"].iter_rows(min_row=2, max_col=14, values_only=True):
            klasor, mahkeme, no, durusma = satir[1], satir[2], satir[3], satir[10]
            if not (klasor or (mahkeme and no)):
                continue
            out.append({"klasor": klasor_kodu(klasor), "mahkeme": str(mahkeme or "").strip(),
                        "no": uo.dosya_no(no) if no else "",
                        "durusma": durusma.isoformat() if isinstance(durusma, (datetime, date)) else None})
        return out, ""
    finally:
        wb.close()


def _kayit_bul(kayitlar: list[dict], mahkeme, no) -> dict | None:
    no = uo.dosya_no(no)
    return next((k for k in kayitlar if k["no"] == no and uo.ayni_dosya(k["mahkeme"], k["no"], mahkeme, no)), None)


def _kayit_ekle(kayitlar: list[dict], mahkeme, no, kaynak: str, portal: str = "", tur: str = "") -> dict:
    k = _kayit_bul(kayitlar, mahkeme, no)
    if k is None:
        k = {"mahkeme": str(mahkeme or "").strip(), "no": uo.dosya_no(no), "portal": portal or "avukat", "tur": tur,
             "kaynaklar": [], "klasor_kesin": {}, "klasor_teyitsiz": [], "cmk": False, "icra": False, "klonlar": [],
             "hareketler": [], "durusmalar": [], "excel_durusma": None, "uygulamada": False}
        kayitlar.append(k)
    if kaynak not in k["kaynaklar"]:
        k["kaynaklar"].append(kaynak)
    if tur and not k["tur"]:
        k["tur"] = tur
    if portal and k["portal"] == "avukat" and portal != "avukat" and kaynak in ("onbellek", "klon"):
        k["portal"] = portal
    return k


def envanter(onbellek: dict, eslesme: list[dict], cmk: list[dict], klonlar: list[dict],
             klasorler: dict[str, list[str]], excel: list[dict]) -> dict:
    """Bürodaki UYAP dosyaları (mahkeme+no) ve UYAP bağı bilinmeyen dava klasörleri."""
    kayitlar: list[dict] = []
    for x in (onbellek.get("dosyalarim") or {}).get("davalar") or []:
        k = _kayit_ekle(kayitlar, x.get("birimAdi"), x.get("dosyaNo"), "onbellek", x.get("portal") or "")
        k["uygulamada"] = True
    for x in (onbellek.get("hareketler") or {}).get("hareketler") or []:
        k = _kayit_ekle(kayitlar, x.get("birimAdi"), x.get("dosyaNo"), "hareket", tur=x.get("dosyaTuru", ""))
        k["hareketler"].append({"tarih": str(x.get("tarih") or "")[:10], "aciklama": x.get("aciklama") or "",
                                "dikkat": uo.dikkat(x.get("aciklama"))})
    for x in (onbellek.get("durusmalar") or {}).get("durusmalar") or []:
        k = _kayit_ekle(kayitlar, x.get("birimAdi"), x.get("dosyaNo"), "durusma")
        k["durusmalar"].append({"tarih": x.get("tarih"), "islem": x.get("islemTuru") or ""})
    for x in klonlar:
        k = _kayit_ekle(kayitlar, x["mahkeme"], x["no"], "klon", x["portal"])
        k["klonlar"].append(x["klon"])
        k["icra"] = k["icra"] or x["icra"]
        if x["klasor"]:
            k["klasor_kesin"].setdefault(x["klasor"], []).append("klon klasörü")
    for e in eslesme:
        k = _kayit_ekle(kayitlar, e["mahkeme"], e["no"], f"eslesme-{e['bolum']}")
        k["cmk"] = k["cmk"] or bool(e.get("cmk"))
        if e.get("klasor") and e["bolum"] == "A":
            k["klasor_kesin"].setdefault(e["klasor"], []).append("eşleşme tablosu A")
        elif e.get("klasor") and e["klasor"] not in k["klasor_teyitsiz"]:
            k["klasor_teyitsiz"].append(e["klasor"])
    for x in excel:
        if x["mahkeme"] and x["no"]:
            k = _kayit_ekle(kayitlar, x["mahkeme"], x["no"], "excel")
            if x["klasor"]:
                k["klasor_kesin"].setdefault(x["klasor"], []).append("Excel")
            if x["durusma"]:
                k["excel_durusma"] = x["durusma"]
    for c in cmk:
        k = _kayit_ekle(kayitlar, c["mahkeme"], c["no"], "cmk-listesi")
        k["cmk"] = True
    cakisan = sorted(kod for kod, adlar in klasorler.items() if len(adlar) > 1)
    for k in kayitlar:
        k["icra"] = k["icra"] or "icra dairesi" in uo.tr(k["mahkeme"])
        kodlar = sorted(k["klasor_kesin"])
        k["buro_kodu"] = kodlar[0] if len(kodlar) == 1 and kodlar[0] not in cakisan else ""
        k["buro_kodu_sorunu"] = ("birden çok kesin klasör: " + ", ".join(kodlar)) if len(kodlar) > 1 else \
            (f"klasör kodu {kodlar[0]} birden çok klasörde" if kodlar and kodlar[0] in cakisan else "")
        k["external_ref"] = uyap_ref(k["mahkeme"], k["no"], k["portal"])
    bagli = {kod for k in kayitlar for kod in k["klasor_kesin"]}
    yalniz_klasor = [{"buro_kodu": kod, "klasor_sayisi": len(adlar), "klasorler": adlar}
                     for kod, adlar in sorted(klasorler.items()) if kod not in bagli]
    return {"dosyalar": kayitlar, "yalniz_klasor": yalniz_klasor, "cakisan_kodlar": cakisan,
            "klasorler": klasorler}


# --- eşleştirme -------------------------------------------------------------------------------------------------

def _aday_numara(dava: dict, dosyalar: list[dict]) -> list[dict]:
    no, mah = dava.get("caseNumber"), dava.get("courtName")
    if not (no and mah) or not re.search(r"\d{4}\s*[/-]\s*\d+", str(no)):
        return []
    return [k for k in dosyalar if k["no"] == uo.dosya_no(no) and uo.ayni_dosya(mah, no, k["mahkeme"], k["no"])]


def eslestir(davalar: list[dict], env: dict) -> dict:
    """isbu davaları ↔ büro dosyaları. Yalnız tek-tek karşılık eşleşir; gerisi belirsiz/eşleşmeyen."""
    dosyalar = env["dosyalar"]
    klasorler = env["klasorler"]
    eslesen, belirsiz, isbu_yok = [], [], []
    for d in davalar:
        ext, kod = (d.get("externalRef") or "").strip(), (d.get("buroKodu") or "").strip()
        aday, yol, neden = [], "", ""
        if ext:
            aday, yol = [k for k in dosyalar if k["external_ref"] == ext], "external_ref"
        if not aday and kod:
            yol = "buro_kodu"
            if kod in env["cakisan_kodlar"]:
                neden = f"büro kodu {kod} birden çok klasörde"
            aday = [k for k in dosyalar if kod in k["klasor_kesin"]]
            if len(aday) > 1:
                daralt = [k for k in aday if k in _aday_numara(d, aday)]
                aday = daralt or aday
            if not aday and kod in klasorler and not neden:
                aday = [{"yalniz_klasor": True, "buro_kodu": kod, "external_ref": "", "klasorler": klasorler[kod]}]
        if not aday and not neden:
            yol = "esas no + mahkeme"
            aday = _aday_numara(d, dosyalar)
            if not aday and d.get("caseNumber") and not d.get("courtName"):
                ayni_no = [k for k in dosyalar if k["no"] == uo.dosya_no(d["caseNumber"])]
                if ayni_no:
                    neden, aday = "mahkeme adı boş; esas no birden çok mahkemede olabilir", ayni_no
        if neden or len(aday) > 1:
            belirsiz.append({"dava": d, "yol": yol, "neden": neden or f"{len(aday)} aday", "adaylar": aday})
        elif aday:
            eslesen.append({"dava": d, "dosya": aday[0], "yol": yol})
        else:
            isbu_yok.append(d)
    # aynı büro dosyasına iki isbu davası bağlanamaz
    sayac: dict[int, int] = {}
    for e in eslesen:
        sayac[id(e["dosya"])] = sayac.get(id(e["dosya"]), 0) + 1
    tekil = [e for e in eslesen if sayac[id(e["dosya"])] == 1]
    for e in eslesen:
        if sayac[id(e["dosya"])] > 1:
            belirsiz.append({"dava": e["dava"], "yol": e["yol"], "neden": "aynı büro dosyası birden çok isbu davasına "
                             "karşılık geliyor", "adaylar": [e["dosya"]]})
    bagli = {id(e["dosya"]) for e in tekil} | {id(a) for b in belirsiz for a in b["adaylar"]}
    buro_yok = [k for k in dosyalar if id(k) not in bagli]
    bagli_kod = {e["dosya"].get("buro_kodu") for e in tekil} | {e["dava"].get("buroKodu") for e in tekil}
    klasor_yok = [y for y in env["yalniz_klasor"] if y["buro_kodu"] not in bagli_kod]
    return {"eslesen": tekil, "belirsiz": belirsiz, "isbu_yok": isbu_yok, "buro_yok": buro_yok,
            "klasor_yok": klasor_yok}


# --- öneri planı ------------------------------------------------------------------------------------------------

def oge(kind: str, external_ref: str, case_id: str, payload: dict, summary: str, reason: str, source: str) -> dict:
    """Şemaya göre temizlenmiş, deterministik request_id'li öneri öğesi. Şema hatasında ValueError."""
    import isbu_ofis
    temiz, hatalar = isbu_sema.payload_temizle(kind, payload)
    if hatalar:
        raise ValueError(f"{kind} {external_ref}: " + "; ".join(hatalar))
    case_ref = {"id": case_id}
    item = {"kind": kind, "external_ref": external_ref[:200], "case_ref": case_ref, "payload": temiz,
            "summary": summary.strip()[:1000], "reason": reason.strip()[:5000], "source": source,
            "request_id": isbu_ofis.request_id(kind, external_ref[:200], temiz, case_ref)}
    sorun = isbu_sema.oge_dogrula(item)
    if sorun:
        raise ValueError(f"{kind} {external_ref}: " + "; ".join(sorun))
    return item


def dava_onerisi(d: dict, k: dict, bugun: date) -> tuple[dict | None, list[str]]:
    """Dava bilgisi + bağlama (case_upsert). Yalnız isbu'da BOŞ alanlar önerilir (CMK hariç: açık overwrite)."""
    if k.get("yalniz_klasor"):
        return None, []
    p, notlar, overwrite = {}, [], []
    if not d.get("buroKodu"):
        if k.get("buro_kodu"):
            p["buro_kodu"] = k["buro_kodu"]
        elif k.get("buro_kodu_sorunu"):
            notlar.append(f"büro kodu önerilmedi: {k['buro_kodu_sorunu']}")
        elif k.get("klasor_teyitsiz"):
            notlar.append("büro kodu önerilmedi: klasör yalnız eşleşme tablosu B (teyit gerekli): "
                          + ", ".join(k["klasor_teyitsiz"]))
    if not d.get("caseNumber") and k.get("no"):
        p["case_number"] = k["no"]
    if not d.get("courtName") and k.get("mahkeme"):
        p["court_name"] = k["mahkeme"]
    if k.get("cmk") and not d.get("isCmkAssignment"):
        p["is_cmk_assignment"] = True
        overwrite.append("is_cmk_assignment")
        notlar.append("CMK listesinde/eşleşme tablosunda CMK; isbu'da CMK işaretsiz (avukat incelemeli)")
    if overwrite:
        p["overwrite"] = overwrite
    ext = d.get("externalRef") or k["external_ref"]
    if not p and d.get("externalRef"):
        return None, notlar
    kaynak = ", ".join(k["kaynaklar"])
    alanlar = ", ".join(x for x in p if x != "overwrite") or "yalnız dış kimlik bağlama"
    item = oge("case_upsert", ext, d["id"], p, f"Dava bilgisi: {alanlar}",
               f"Büro kaynakları ({kaynak}) {bugun:%d.%m.%Y}; eşleşme esas no + mahkeme veya büro kodu ile tekil. "
               + " ".join(notlar), "uyap")
    return item, notlar


def durusma_onerileri(d: dict, k: dict, detay: dict | None, bugun: date) -> list[dict]:
    out, mevcut = [], (detay or {}).get("hearings") or []
    for x in k.get("durusmalar") or []:
        iso = tr_iso(x.get("tarih"))
        if not iso or datetime.fromisoformat(iso).date() < bugun:
            continue
        if any(ayni_dakika(h.get("hearingDate"), iso) for h in mevcut):
            continue
        ext = f"{k['external_ref']}:durusma:{iso[:16]}"
        out.append(oge("hearing_upsert", ext, d["id"], {"hearing_date": iso, "notes": x.get("islem") or ""},
                       f"UYAP: duruşma {iso[:10]} {iso[11:16]}",
                       f"UYAP önbelleği duruşma listesi ({k['mahkeme']} {k['no']}).", "uyap"))
    return out


def _kalkan_baslik(s: dict) -> str:
    return str(s.get("islem") or s.get("satir_id") or "Süre").strip()[:440]


def sure_onerileri(d: dict, satirlar: list[dict], detay: dict | None) -> tuple[list[dict], list[str]]:
    """Süre Kalkanı satırları -> deadline_upsert. Kesin: due_date; koşullu/tarihsiz: conditional, due_date boş."""
    out, notlar = [], []
    acik = {t.get("externalRef"): t for t in (detay or {}).get("openDeadlines") or [] if t.get("externalRef")}
    for s in satirlar:
        if not isinstance(s, dict) or s.get("durum") == "arşiv" or not s.get("satir_id"):
            continue
        ext = f"kalkan:{s['satir_id']}"
        p = {"title": _kalkan_baslik(s), "description": f"Kaynak: KALKAN#{s['satir_id']}",
             "legal_basis": next((str(s[a]) for a in ("dayanak", "kanuni_dayanak", "legal_basis") if s.get(a)), "")}
        kos = s.get("kosullu_hesap") if isinstance(s.get("kosullu_hesap"), dict) else {}
        if isbu_sema.gun_mu(s.get("hukuki_son_gun")):
            p["due_date"] = s["hukuki_son_gun"]
            ozet = f"Süre (kesin): {p['title']} son gün {s['hukuki_son_gun']}"
        elif kos.get("tarih"):
            p.update(conditional=True, condition=f"Başlangıç teyidi bekliyor ({kos.get('baslangic_durumu') or '?'}); "
                     f"koşullu hesap {kos['tarih']} — kesin son gün değil")
            ozet = f"Süre (koşullu): {p['title']}"
        else:
            p.update(conditional=True, condition=f"Tarihsiz Kalkan satırı (durum: {s.get('durum') or '?'})")
            ozet = f"Süre (tarihsiz): {p['title']}"
        var = acik.get(ext)
        if var and p.get("due_date") and str(var.get("dueDate") or "")[:10] == p["due_date"]:
            notlar.append(f"{ext}: isbu'da güncel")
            continue
        out.append(oge("deadline_upsert", ext, d["id"], p, ozet, f"Süre Kalkanı KALKAN#{s['satir_id']}", "kalkan"))
    return out, notlar


def icra_onerisi(d: dict, ozet: dict, detay: dict | None, k: dict) -> dict | None:
    veri = {a: ozet.get(a) for a in isbu_sema.ENFORCEMENT_KEYS if isinstance(ozet.get(a), str) and ozet[a].strip()}
    mevcut = (((detay or {}).get("enforcement") or {}).get("data")) or {}
    degisen = {a: v for a, v in veri.items() if mevcut.get(a) != v}
    if not degisen:
        return None
    return oge("enforcement_upsert", f"icra:{k['external_ref']}", d["id"],
               {"data": degisen, "source_ref": f"yargipro-klon:{norm_mahkeme(k['mahkeme'])}|{k['no']}"[:200]},
               f"İcra özeti: {len(degisen)} alan", "icra_ozet: Yargı PRO klonundaki dosya bilgileri, safahat ve sorgu "
               "evrakı (UYAP'a yeni sorgu yapılmadı). 'Sorgu evrakı yok' mal varlığı yok demek değildir.", "icra_ozet")


def sonraki_adim_onerisi(d: dict, k: dict, bugun: date) -> dict | None:
    """Sırada işi olmayan/atıl dosyada son 30 günün UYAP 'Dikkat' olayından bekleme nedeni (tarih uydurulmaz)."""
    if not set(d.get("flags") or []) & {"atil", "sirada_is_yok"}:
        return None
    if ((d.get("workspace") or {}).get("waitingFor") or "").strip():
        return None
    sinir = (bugun - timedelta(days=30)).isoformat()
    olaylar = sorted((h for h in k.get("hareketler") or [] if h["dikkat"] and h["tarih"] >= sinir),
                     key=lambda h: h["tarih"], reverse=True)
    if not olaylar:
        return None
    h = olaylar[0]
    bekleme = f"{h['dikkat']} (UYAP {h['tarih']}: {h['aciklama']})"[:500]
    return oge("next_step", f"{k['external_ref']}:sonraki-adim", d["id"], {"target": "workspace", "waiting_for": bekleme},
               f"Sıradaki adım: {h['dikkat']}", "isbu nabzı: sırada iş yok/atıl; bekleme nedeni UYAP hareketinden "
               "(uyap_ortak.dikkat kuralı). Kontrol tarihi önerilmedi (kaynak yok).", "uyap")


def plan_kur(sonuc: dict, detaylar: dict, kalkan: dict, icra_ozetleri: dict, bugun: date) -> dict:
    """Eşleşen her dosya için öneri öğeleri + önerilmeyenlerin gerekçesi.
    detaylar: {case_id: GET /cases/:id}, kalkan: {buro_kodu: [satır]|None}, icra_ozetleri: {external_ref: özet}."""
    items, kayit = [], []
    for e in sonuc["eslesen"]:
        d, k = e["dava"], e["dosya"]
        r = {"case_id": d["id"], "buro_kodu": d.get("buroKodu") or k.get("buro_kodu") or "", "mahkeme":
             k.get("mahkeme", ""), "no": k.get("no", ""), "oneriler": [], "kaynak_yok": [], "notlar": []}
        if d.get("status") == "closed":
            r["notlar"].append("isbu'da kapalı dava: öneri üretilmedi")
            kayit.append(r)
            continue
        detay = detaylar.get(d["id"])
        it, notlar = dava_onerisi(d, k, bugun)
        r["notlar"] += notlar
        yeni = [it] if it else []
        if not k.get("yalniz_klasor"):
            dur = durusma_onerileri(d, k, detay, bugun)
            yeni += dur
            if not k.get("durusmalar"):
                r["kaynak_yok"].append("duruşma: UYAP önbelleğinde duruşma listesi yok")
            if k.get("excel_durusma") and not dur:
                r["notlar"].append(f"Excel'de duruşma {k['excel_durusma'][:10]} (saat/UYAP teyidi yok; önerilmedi)")
        kod = r["buro_kodu"]
        satirlar = kalkan.get(kod) if kod else None
        if satirlar:
            s_items, s_not = sure_onerileri(d, satirlar, detay)
            yeni += s_items
            r["notlar"] += s_not
        else:
            r["kaynak_yok"].append("süre: Süre Kalkanı kaydı yok" if kod else "süre: büro kodu yok, Kalkan aranamadı")
            olay = [h for h in k.get("hareketler") or [] if h["dikkat"].startswith(SURE_DIKKAT)]
            if olay:
                r["notlar"].append("süre doğurabilecek UYAP olayı var, Kalkan'a hesaplanmadı: "
                                   + "; ".join(f"{h['tarih']} {h['dikkat']}" for h in olay[:3]))
        if k.get("icra"):
            oz = icra_ozetleri.get(k.get("external_ref"))
            if oz:
                it = icra_onerisi(d, oz, detay, k)
                if it:
                    yeni.append(it)
                else:
                    r["notlar"].append("icra özeti isbu'da güncel")
            else:
                r["kaynak_yok"].append("icra özeti: Yargı PRO icra klonu yok")
        it = sonraki_adim_onerisi(d, k, bugun)
        if it:
            yeni.append(it)
        elif set(d.get("flags") or []) & {"atil", "sirada_is_yok"}:
            r["kaynak_yok"].append("sıradaki adım: son 30 günde UYAP 'Dikkat' olayı yok")
        r["oneriler"] = [{"kind": x["kind"], "summary": x["summary"], "external_ref": x["external_ref"]} for x in yeni]
        items += yeni
        kayit.append(r)
    return {"items": items, "dosyalar": kayit}


# --- rapor ------------------------------------------------------------------------------------------------------

def _sayim(items: list[dict]) -> dict:
    out: dict = {}
    for x in items:
        out[x["kind"]] = out.get(x["kind"], 0) + 1
    return dict(sorted(out.items()))


def _baslik(d: dict) -> str:
    return re.sub(r"[|\s]+", " ", str(d.get("title") or "")).strip()[:60]


def _dava_satiri(d: dict) -> str:
    return (f"| {d.get('id', '')[:8]} | {_baslik(d)} | {d.get('buroKodu') or '—'} | {d.get('courtName') or '—'} | "
            f"{d.get('caseNumber') or '—'} | {d.get('status') or ''}{' · CMK' if d.get('isCmkAssignment') else ''} |")


def md_eslesme(sonuc: dict, env: dict, uyarilar: list[str], bugun: date) -> str:
    s = [f"# isbu-ofis ↔ büro eşleştirmesi · {bugun:%d.%m.%Y}", "",
         "Kuru çalışma. UYAP'a gidilmedi; isbu-ofis'ten yalnız okundu. Belirsiz olan zorla eşleştirilmedi.", "",
         f"- Eşleşen: **{len(sonuc['eslesen'])}** · Belirsiz: **{len(sonuc['belirsiz'])}** · isbu-ofis'te olup büroda "
         f"karşılığı olmayan: **{len(sonuc['isbu_yok'])}** · Büroda olup isbu-ofis'te olmayan UYAP dosyası: "
         f"**{len(sonuc['buro_yok'])}** · isbu'ya bağlanmamış dava klasörü: **{len(sonuc['klasor_yok'])}**",
         f"- Klasör kodu çakışması: {', '.join(env['cakisan_kodlar']) or 'yok'}", ""]
    if uyarilar:
        s += ["**Okunamayan / eksik kaynak:** " + "; ".join(uyarilar), ""]
    s += ["## Eşleşen", "", "| isbu id | Başlık | Büro kodu | Mahkeme (büro) | Esas no | Yol |", "|---|---|---|---|---|---|"]
    s += [f"| {e['dava']['id'][:8]} | {_baslik(e['dava'])} | {e['dava'].get('buroKodu') or e['dosya'].get('buro_kodu') or '—'} | "
          f"{e['dosya'].get('mahkeme') or '(yalnız klasör)'} | {e['dosya'].get('no') or '—'} | {e['yol']} |"
          for e in sonuc["eslesen"]] or ["| — | — | — | — | — | — |"]
    s += ["", "## Belirsiz (öneri üretilmez; avukat teyidi)", "", "| isbu id | Başlık | Büro kodu | Mahkeme | Esas no | Neden |",
          "|---|---|---|---|---|---|"]
    s += [f"| {b['dava']['id'][:8]} | {_baslik(b['dava'])} | {b['dava'].get('buroKodu') or '—'} | {b['dava'].get('courtName') or '—'} | "
          f"{b['dava'].get('caseNumber') or '—'} | {b['neden']} |" for b in sonuc["belirsiz"]] or ["| — | — | — | — | — | — |"]
    s += ["", "## isbu-ofis'te olup büro kaynaklarında karşılığı olmayan", "",
          "| isbu id | Başlık | Büro kodu | Mahkeme | Esas no | Durum |", "|---|---|---|---|---|---|"]
    s += [_dava_satiri(d) for d in sonuc["isbu_yok"]] or ["| — | — | — | — | — | — |"]
    gruplar = {"CMK": [k for k in sonuc["buro_yok"] if k["cmk"]],
               "İcra": [k for k in sonuc["buro_yok"] if k["icra"] and not k["cmk"]],
               "Dava": [k for k in sonuc["buro_yok"] if not k["icra"] and not k["cmk"]]}
    s += ["", "## Büroda olup isbu-ofis'te olmayan UYAP dosyaları", ""]
    for ad, liste in gruplar.items():
        s += [f"### {ad} ({len(liste)})", "", "| Mahkeme / Daire | Dosya no | Klasör | Kaynak |", "|---|---|---|---|"]
        s += [f"| {k['mahkeme']} | {k['no']} | {k.get('buro_kodu') or ', '.join(k['klasor_teyitsiz']) or '—'} | "
              f"{', '.join(k['kaynaklar'])} |" for k in liste] or ["| — | — | — | — |"]
        s.append("")
    s += ["## isbu-ofis'e bağlanmamış dava klasörleri (UYAP bağı bilinmiyor)", ""]
    s += [f"- {y['buro_kodu']}" + (f" ({y['klasor_sayisi']} klasör — kod çakışması)" if y["klasor_sayisi"] > 1 else "")
          for y in sonuc["klasor_yok"]] or ["- yok"]
    return "\n".join(s) + "\n"


def md_plan(plan: dict, sonuc: dict, bugun: date, gonderildi: bool = False) -> str:
    sayim = _sayim(plan["items"])
    kaynak_yok: dict = {}
    for r in plan["dosyalar"]:
        for x in r["kaynak_yok"]:
            anahtar = x.split(":")[0]
            kaynak_yok[anahtar] = kaynak_yok.get(anahtar, 0) + 1
    s = [f"# isbu-ofis öneri planı · {bugun:%d.%m.%Y}", "",
         ("**GÖNDERİLDİ** (avukat onayıyla)." if gonderildi else
          "**KURU ÇALIŞMA — hiçbir öneri gönderilmedi.** Gönderim yalnız avukat onayıyla `plan --gonder`."), "",
         f"- Eşleşen dosya: {len(sonuc['eslesen'])} · önerilecek öğe: **{len(plan['items'])}**",
         "- Tür bazında: " + (", ".join(f"{k} {v}" for k, v in sayim.items()) or "yok"),
         "- Kaynak olmadığı için önerilmeyen (dosya sayısı): "
         + (", ".join(f"{k} {v}" for k, v in sorted(kaynak_yok.items())) or "yok"),
         "- Süre önerisi yalnız Süre Kalkanı satırından üretilir; tarih uydurulmaz.", ""]
    for r in plan["dosyalar"]:
        s += [f"## {r['buro_kodu'] or '—'} · {r['mahkeme'] or '(yalnız klasör)'} {r['no']} · isbu {r['case_id'][:8]}", ""]
        s += [f"- ÖNERİ `{o['kind']}`: {o['summary']} (`{o['external_ref']}`)" for o in r["oneriler"]] or ["- öneri yok"]
        s += [f"- kaynak yok — {x}" for x in r["kaynak_yok"]]
        s += [f"- not: {x}" for x in r["notlar"]]
        s.append("")
    return "\n".join(s) + "\n"


# --- yerel kaynak + isbu okuması (CLI) ---------------------------------------------------------------------------

def yerel_kaynaklar() -> tuple[dict, list[str]]:
    from paths import aktif_davalar, data_root
    import excel_takip
    kok = aktif_davalar()
    uyarilar = []
    onbellek = uo.onbellek_oku()
    if not onbellek:
        uyarilar.append("UYAP önbelleği yok")
    else:
        for anahtar, neden in (onbellek.get("hatalar") or {}).items():
            uyarilar.append(f"önbellek {onbellek.get('tarih')} · {anahtar}: okunmadı ({str(neden)[:60]})")
    excel, uyari = excel_satirlari(data_root() / excel_takip.DOSYA_ADI)
    if uyari:
        uyarilar.append(uyari)
    env = envanter(onbellek, uo.eslesme_satirlari(), uo.cmk_listesi(), klon_tara(kok), dava_klasorleri(kok), excel)
    return env, uyarilar


def kalkan_oku(sonuc: dict, klasorler: dict) -> tuple[dict, list[str]]:
    import sure_kalkani
    from paths import aktif_davalar
    out, uyari = {}, []
    for e in sonuc["eslesen"]:
        kod = e["dava"].get("buroKodu") or e["dosya"].get("buro_kodu")
        if not kod or kod in out or len(klasorler.get(kod, [])) != 1:
            continue
        try:
            depo = sure_kalkani.yukle(aktif_davalar() / klasorler[kod][0])
            out[kod] = depo["satirlar"] or None
        except sure_kalkani.OKUMA_HATALARI as hata:
            out[kod] = None
            uyari.append(f"Kalkan okunamadı ({kod}): {type(hata).__name__}")
    return out, uyari


def icra_oku(sonuc: dict, bugun: date) -> tuple[dict, list[str]]:
    import icra_ozet
    out, uyari = {}, []
    for e in sonuc["eslesen"]:
        k = e["dosya"]
        if not k.get("icra") or not k.get("klonlar") or e["dava"].get("status") == "closed":
            continue
        try:
            out[k["external_ref"]] = icra_ozet.dosya_ozeti(Path(k["klonlar"][0]), bugun)
        except (OSError, ValueError, KeyError, TypeError) as hata:
            uyari.append(f"icra özeti çıkarılamadı ({k['no']}): {type(hata).__name__}")
    return out, uyari


def _yaz(klasor: Path, ad: str, icerik) -> Path:
    klasor.mkdir(parents=True, exist_ok=True)
    yol = klasor / ad
    metin = icerik if isinstance(icerik, str) else json.dumps(icerik, ensure_ascii=False, indent=1, default=str)
    yol.write_text(metin, encoding="utf-8")
    return yol


def _json_guvenli(env_dosyasi: dict) -> dict:
    return {k: v for k, v in env_dosyasi.items() if k not in ("klonlar",)}


def calistir(komut: str, cikti: Path, gonder: bool = False, api=None, bugun: date | None = None) -> dict:
    import isbu_ofis
    bugun = bugun or date.today()
    api = api or isbu_ofis.IsbuOfis()
    env, uyarilar = yerel_kaynaklar()
    try:
        api.saglik()
    except isbu_ofis.IsbuHata:
        pass  # uyandırma denemesi; asıl okuma kendi yeniden denemesini yapar
    davalar = api.cases(include_closed=True)
    sonuc = eslestir(davalar, env)
    _yaz(cikti, "eslesme.md", md_eslesme(sonuc, env, uyarilar, bugun))
    _yaz(cikti, "eslesme.json", {
        "tarih": bugun.isoformat(), "uyarilar": uyarilar, "cakisan_kodlar": env["cakisan_kodlar"],
        "eslesen": [{"case_id": e["dava"]["id"], "yol": e["yol"], "dosya": _json_guvenli(e["dosya"])} for e in sonuc["eslesen"]],
        "belirsiz": [{"case_id": b["dava"]["id"], "neden": b["neden"], "aday": len(b["adaylar"])} for b in sonuc["belirsiz"]],
        "isbu_yok": [d["id"] for d in sonuc["isbu_yok"]],
        "buro_yok": [{"mahkeme": k["mahkeme"], "no": k["no"], "cmk": k["cmk"], "icra": k["icra"]} for k in sonuc["buro_yok"]],
        "klasor_yok": [y["buro_kodu"] for y in sonuc["klasor_yok"]]})
    ozet = {"eslesen": len(sonuc["eslesen"]), "belirsiz": len(sonuc["belirsiz"]), "isbu_yok": len(sonuc["isbu_yok"]),
            "buro_yok": len(sonuc["buro_yok"]), "klasor_yok": len(sonuc["klasor_yok"]), "uyarilar": uyarilar}
    if komut == "eslestir":
        return ozet
    acik = [e for e in sonuc["eslesen"] if e["dava"].get("status") != "closed"]
    detaylar = {e["dava"]["id"]: api.case(e["dava"]["id"]) for e in acik}
    kalkan, u1 = kalkan_oku(sonuc, env["klasorler"])
    icra, u2 = icra_oku(sonuc, bugun)
    plan = plan_kur(sonuc, detaylar, kalkan, icra, bugun)
    ozet.update(oge=len(plan["items"]), tur=_sayim(plan["items"]), uyarilar=uyarilar + u1 + u2)
    gonderildi = None
    if gonder and plan["items"]:
        gonderildi = api.post_proposals(plan["items"], gonder=True)
        ozet["gonderim"] = [{"http": g.get("http"), "counts": g.get("counts")} for g in gonderildi]
    _yaz(cikti, "plan.md", md_plan(plan, sonuc, bugun, gonderildi is not None))
    _yaz(cikti, "plan.json", {"tarih": bugun.isoformat(), "kuru": gonderildi is None, "sayim": _sayim(plan["items"]),
                              "items": plan["items"], "dosyalar": plan["dosyalar"], "gonderim": gonderildi})
    return ozet


def main(argv: list[str]) -> int:
    for akim in (sys.stdout, sys.stderr):
        try:
            akim.reconfigure(encoding="utf-8")
        except (AttributeError, ValueError):
            pass
    ap = argparse.ArgumentParser(prog="isbu_esitle.py")
    alt = ap.add_subparsers(dest="komut", required=True)
    for ad in ("eslestir", "plan"):
        p = alt.add_parser(ad)
        p.add_argument("--cikti", default="")
        if ad == "plan":
            p.add_argument("--gonder", action="store_true", help="önerileri GÖNDERİR (yalnız avukat onayıyla)")
    args = ap.parse_args(argv)
    cikti = Path(args.cikti) if args.cikti else REPO_ROOT / "tmp" / "isbu-ofis" / date.today().isoformat()
    import isbu_ofis
    try:
        ozet = calistir(args.komut, cikti, gonder=getattr(args, "gonder", False))
    except (isbu_ofis.IsbuHata, OSError, ValueError, RuntimeError) as hata:
        print(f"HATA: {hata}", file=sys.stderr)
        return 1
    print(json.dumps(ozet, ensure_ascii=False, indent=1))
    print(f"Rapor: {cikti}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main(sys.argv[1:]))
