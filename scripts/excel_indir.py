#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""excel_indir.py — OneDrive'daki canlı iş takip Excel'ini (VEGA HUKUK .xlsx) yerel kopyaya indirir.

Avukat 22.10.2026: Excel OneDrive üzerinden güncel tutuluyor; masaüstü kopyası eski. Bu betik canlı nüshayı
`config/paths.json` → `buro_excel` yoluna yazar (önce yedek `<ad>_eski_<tarih-saat>.xlsx`).

İki yol:
1. PAYLAŞIM LİNKİ (`--link` ya da config `buro_excel_onedrive.paylasim`): OneDrive'da "Bağlantıya sahip herkes
   görüntüleyebilir" linki. OneDrive shares API (`u!` + base64url(link)) ile oturum açmadan indirilir.
   Avukatın doc2.aspx görüntüleme linki bu yolla ÇALIŞMAZ (401 unauthenticated).
2. DOSYA (`--dosya <xlsx>`): Chrome'un (OneDrive oturumu açık) `download.aspx?UniqueId=…` ile indirdiği dosya;
   doğrulanıp yerine konur. (Chrome: config `buro_excel_onedrive.download` adresi.)

Kullanım:
  python scripts/excel_indir.py [--link <paylaşım linki>] [--dosya <indirilen xlsx>] [--kuru]
"""
from __future__ import annotations

import argparse
import base64
import json
import shutil
import sys
import urllib.error
import urllib.request
from datetime import datetime
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))
from paths import REPO_ROOT, buro_excel  # noqa: E402

XLSX_MAGIC = b"PK\x03\x04"


def paylasim_indir(link: str) -> bytes:
    tok = "u!" + base64.urlsafe_b64encode(link.encode("utf-8")).decode("ascii").rstrip("=")
    url = f"https://api.onedrive.com/v1.0/shares/{tok}/root/content"
    try:
        with urllib.request.urlopen(urllib.request.Request(url, headers={"User-Agent": "Mozilla/5.0"}), timeout=60) as r:
            return r.read()
    except urllib.error.HTTPError as h:
        raise RuntimeError(f"OneDrive indirme HTTP {h.code}: paylaşım linki 'bağlantıya sahip herkes' olmalı "
                           f"(görüntüleme linki yetmez)") from None


def dogrula(veri: bytes) -> dict:
    if not veri.startswith(XLSX_MAGIC):
        raise RuntimeError("indirilen içerik .xlsx değil (giriş sayfası olabilir)")
    import io
    import openpyxl
    wb = openpyxl.load_workbook(io.BytesIO(veri), read_only=True, data_only=True)
    try:
        if "DOSYALAR" not in wb.sheetnames:
            raise RuntimeError(f"DOSYALAR sayfası yok: {wb.sheetnames}")
        ws = wb["DOSYALAR"]
        satir = sum(1 for r in ws.iter_rows(min_row=5, max_col=3, values_only=True) if r[2])
        return {"sayfalar": wb.sheetnames, "dosya_satiri": satir, "ozet": ws.cell(2, 1).value}
    finally:
        wb.close()


def yerlestir(veri: bytes, hedef: Path, kuru: bool) -> Path | None:
    bilgi = dogrula(veri)
    print(f"doğrulandı: {bilgi}")
    if kuru:
        return None
    if hedef.exists() and hedef.read_bytes() == veri:
        print("değişiklik yok")
        return hedef
    if hedef.exists():
        yedek = hedef.with_name(f"{hedef.stem}_eski_{datetime.now():%Y-%m-%d_%H%M}{hedef.suffix}")
        shutil.copyfile(hedef, yedek)
        print(f"yedek: {yedek.name}")
    hedef.write_bytes(veri)
    print(f"yazıldı: {hedef}")
    return hedef


def main(argv: list[str]) -> int:
    ap = argparse.ArgumentParser(prog="excel_indir.py")
    ap.add_argument("--link", help="OneDrive paylaşım linki (bağlantıya sahip herkes)")
    ap.add_argument("--dosya", help="Chrome ile indirilmiş .xlsx (doğrulanıp yerine konur)")
    ap.add_argument("--kuru", action="store_true")
    a = ap.parse_args(argv)
    from paths import _load_cfg  # paths.json + gitignore'daki paths.local.json (link yalnız orada)
    link = a.link or (_load_cfg().get("buro_excel_onedrive") or {}).get("paylasim")
    if a.dosya:
        veri = Path(a.dosya).read_bytes()
    elif link:
        veri = paylasim_indir(link)
    else:
        print("Kaynak yok: --dosya ya da --link ver (ya da config buro_excel_onedrive.paylasim).", file=sys.stderr)
        return 2
    try:
        yerlestir(veri, buro_excel(), a.kuru)
    except RuntimeError as h:
        print(f"[HATA] {h}", file=sys.stderr)
        return 1
    return 0


if __name__ == "__main__":
    sys.exit(main(sys.argv[1:]))
